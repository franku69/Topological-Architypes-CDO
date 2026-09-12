from pathlib import Path
import sys,json
import pandas as pd
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
root=Path(__file__).resolve().parents[1];sys.path.insert(0,str(root/'code'))
from model import reference_specs
R=root/'results';F=root/'figures'
findings=json.loads((R/'computed_findings.json').read_text());log=json.loads((R/'execution_log.json').read_text());ptlog=json.loads((R/'paired_timing_execution.json').read_text())
main=pd.read_csv(R/'summary_primary.csv');timing=pd.read_csv(R/'paired_timing_summary.csv');storage=pd.read_csv(R/'storage_paired.csv');desc=pd.read_csv(R/'topology_descriptors.csv');area=pd.read_csv(R/'area_weight_paired.csv')
doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
sec.top_margin=Inches(.68);sec.bottom_margin=Inches(.65);sec.left_margin=Inches(.68);sec.right_margin=Inches(.68)
sec.header_distance=Inches(.28);sec.footer_distance=Inches(.28)
style=doc.styles['Normal'];style.font.name='Calibri';style.font.size=Pt(10.5);style.paragraph_format.space_after=Pt(6);style.paragraph_format.line_spacing=1.1
for hn,size in [('Title',25),('Heading 1',18),('Heading 2',12)]:
 st=doc.styles[hn];st.font.name='Calibri';st.font.size=Pt(size);st.font.color.rgb=RGBColor.from_string('183648');st.paragraph_format.space_after=Pt(8)
header=sec.header.paragraphs[0];header.text='LAPASAN-MOTIVATED STUDY  |  SIX-CASE TOPOLOGY EXPERIMENT';header.style='Caption';header.runs[0].font.size=Pt(8)
f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=f.add_run('Computational supplement | Page ');r.font.size=Pt(8)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)

def p(t,bold=False,style=None):
 a=doc.add_paragraph(style=style)
 r=a.add_run(t);r.bold=bold
 return a

def h(t,level=1):doc.add_heading(t,level=level)
def page(t):doc.add_page_break();h(t)
def table(head,rows,widths=None):
 t=doc.add_table(rows=1,cols=len(head));t.style='Light Shading Accent 1'
 for c,x in zip(t.rows[0].cells,head):c.text=str(x)
 rep=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(rep)
 for row in rows:
  cells=t.add_row().cells
  for c,x in zip(cells,row):c.text=str(x)
 for row in t.rows:
  trpr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
  for j,c in enumerate(row.cells):
   if widths:c.width=Inches(widths[j])
   for para in c.paragraphs:
    para.paragraph_format.space_after=Pt(4);para.paragraph_format.line_spacing=1.
    for run in para.runs:run.font.size=Pt(9)
 return t

def fig(name,caption,width=6.8):
 doc.add_picture(str(F/(name+'.png')),width=Inches(width));p(caption,style='Caption')
def pct(v):return f'{100*v:.2f}%'
def pp(v):return f'{100*v:+.2f} pp'

h('Six-Case Hydro-Topological\nStress Simulation',0)
p('CDO and five Southeast Asian analogues',True)
p('A document-constrained computational comparison motivated by flooding concerns in Barangay Lapasan. Executed model results, methods, counterexamples and reproducibility record.')
h('Executive finding',2)
p('The simulations support a conditional arrival-concentration mechanism, not an inevitable or uniquely CDO-specific failure pattern. With the same input volumes, pulse shapes and link parameters, simultaneous source starts increased ensemble-mean downstream spill in all six coarse models. However, individual scenarios reversed the effect, and changing storage-release functions or uncertain junction order changed the results materially.')
table(['Computed result','Value / interpretation'],[
('Paired source-timing effect','Mean spill increased by 6.76–13.21 percentage points across the six neutral-storage receiving probes.'),
('CDO arrangement versus matched alternatives','Reference mean spill: 18.39%. Three balanced alternatives averaged 20.49%. The CDO pattern was not invariably worst.'),
('Storage-function sensitivity','Strong synthetic storage reduced mean receiver spill from 22.05% to 14.08% in Manado and from 17.18% to 7.30% in Huong.'),
('Unspecified Kuranji junction order','105 labelled binary refinements produced 19.31–23.03% mean spill on a fixed subset, versus 24.18% for the collapsed confluence-zone abstraction.'),
('Equal-budget outlet test','Two routes did not improve spill under a shared service reduction; they retained service when only one individual route was blocked.')],[2.1,4.5])
h('What was actually executed',2)
p(f"{log['completed_routes']+ptlog['additional_routes']:,} main/control/paired graph-routing evaluations and {log['probe_evaluations']+log['outlet_counterfactual_evaluations']+ptlog['additional_queue_evaluations']:,} finite-buffer capacity evaluations. Many capacity evaluations reuse the same routed input. These are not independent storms, independent geographic samples or annual flood probabilities.")
p('Status: six report-derived coarse graphs with synthetic parameters. No GIS drainage extraction, terrain flood spreading, water-depth prediction, historical-event calibration or Lapasan inundation mapping was performed.',True)
p('Source basis: six uploaded technical reports, denoted D1–D6. Their underlying external references were not independently verified in this execution. Earlier synthetic plots inside the reports were not treated as observations or calibration targets.')

page('1. Data audit and graph-construction rules')
p('The source reports constrain named branches, broad connections, storage/control placement and receiving boundaries. They are not complete machine-readable river networks. The present dataset therefore contains an explicit, reproducible coarse graph for each report, not an unobserved fine drainage inventory. Graphs and per-edge evidence labels are archived in inputs/.')
table(['Case','Preserved source structure','Unresolved or assumed'],[
('CDO [D1, pp. 2–3, 11–12]','Upper CDO + Tumalaong; then Bubunawan; then lower receiving corridor.','Connected lower laterals aggregated. Source-area partition unknown. Lapasan link not established.'),
('Manado [D2, pp. 3, 5, 12]','Lake/control-affected Tondano and separate Tikala; lower urban receiver.','Conceptual control sequence; synthetic release times; no real gate operations.'),
('Kuranji [D3, pp. 1–3, 7, 13]','Five named source groups enter a middle confluence network.','Complete merge order unspecified. Zone collapse and binary alternatives kept separate.'),
('Huong [D4, pp. 3, 10–12]','Ta/Huu merge at Tuan, pass Hue, then receive Bo at Sinh; lagoon and two inlets.','Residual entry at Sinh provisional. Distributaries aggregated. Outlet split assumed.'),
('Kadamaian [D5, pp. 7, 11]','Upper Kadamaian + Panataran; lower/intermediate residual added separately.','Lowland/coastal boundary aggregated; actual parallel threads not reconstructed.'),
('Hat Yai [D6, pp. 3, 10, 21]','Separate major inflows; urban receiving zone; U-Tapao and Ror.1 routes to common lake.','Entry zone not a six-way point junction. 50:50 route split is a neutral assumption.')],[1.45,2.55,2.6])
h('No silent conversion of context into geometry',2)
p('The Kitanglad-facing headwaters → Libona–Baungon–Talakag → Cabula–Macahambus → urban CDO → Macajalar Bay narrative is retained as regional context. Those place names are not converted into an invented sequence of surveyed graph vertices. D1 places the river mouth between Bonbon–Kauswagan and Macabalan–Puntod (p. 8); it does not establish a direct model edge to Barangay Lapasan.')
p('Terrain slope, rainfall normals, historical discharge, event inundation areas and engineering capacities listed in the reports are not mixed into the normalized baseline. Source-area weighting is used only in a separate sensitivity layer where a consistent partition is supplied.')

page('2. Mathematical model and interpretation')
h('Directed routing with source tracking',2)
p('Let G = (V,E) be a directed acyclic coarse network. Each named source supplies its own non-negative pulse. Total event input is fixed at 100 units. Contributions travel along the permitted edges and merge only at represented junctions or explicitly labelled receiving-zone aggregates. Every source is tracked separately through the calculation.')
p('At observation node v:  q_v(t) = sum_s q_s,v(t).',True)
p('Logical river links delay input by one tick in the baseline, or by 1–4 ticks in heterogeneous-delay scenarios. A linear release-memory filter may act on each logical link. Named storage controls have additional release-memory functions in the storage sensitivity. Zero-delay connector edges around control boxes prevent extra travel steps from being created merely by drawing additional boxes.')
p('For a memory time tau > 0: f = 1 − exp(−1/tau); release_t = f × (stored_(t−1) + input_t); stored_t = stored_(t−1) + input_t − release_t. For tau = 0, the node passes its input immediately.',True)
p('This is a conservative storage/release rule with an unlimited routing store. It is not a measured stage–storage curve, flood-control rule curve, gate schedule or reservoir-safety model. The storage-node values are experimental attributes, not pure connectivity.')
h('Finite-buffer overflow diagnostic',2)
p('Each reported probe is run independently. For incoming volume q_t, service capacity C per tick, and buffer B:')
p('available_t = buffer_(t−1) + q_t\nreleased_t = min(C, available_t)\nspill_t = max(0, available_t − released_t − B)\nbuffer_t = available_t − released_t − spill_t',True)
p('All buffers start empty. Released volume + spilled volume + final buffer volume equals incoming volume. Overflow means positive spill in this defined discrete model. It is not centimetres of water, a flooded polygon, channel overtopping verified in the field, or a municipal flood probability.')
h('Independent probes are not a coupled hydraulic model',2)
p('A Hue-city probe and a Huong-lagoon probe are separate diagnostic experiments. Spill at the first is not subtracted from the second, and their spills must never be added as basin-wide damage or total inundation. The model has no momentum, water surface, tide propagation, reverse flow or floodplain spreading. Shared downstream restriction is tested as a capacity multiplier, not a computed lake level.')
p('For identical source pulse p and unit-delay lossless routing, q_v(t) = sum_k W_v(k) p(t−k), where W_v(k) is the source-weighted count at graph distance k. Thus arrival-path concentration, not branch count alone, controls the baseline response. Different graphs with the same weighted path response can be indistinguishable at one outlet.')

page('3. Scenario design, controls and denominators')
table(['Factor','Executed settings'],[
('Total source input','100 abstract units per routed event; kept fixed after selecting active sources.'),
('Pulse duration and shape','2, 8 or 24 ticks; constant block or triangular pulse.'),
('Source starting-time spread','0, 8 or 24 ticks; random source onsets within the selected spread.'),
('Source weights','Equal or lognormal random weights, normalized to total input.'),
('Source activation','All sources, or the highest random scores selecting ceil(0.6 × source count).'),
('Link delays','Unit delay, or deterministic seeded delays of 1–4 ticks.'),
('Ordinary link memory','Release time 0 or 2 ticks.'),
('Realization index','Four indexed realizations per factor cell. Global seed 260911.'),
('Standalone probe settings','Capacities 2, 5, 10, 20 units/tick; buffers 0, 10, 30 units: 12 combinations.'),
('Named storage functions','Neutral, tau = 2, tau = 8. Manado lake uses tau and dam uses tau/2; Huong reservoirs each use tau.'),
('Time horizon','512 abstract ticks; selected runs repeated at 1,024 ticks for closeout checks.')],[2.,4.6])
p('This factorial grammar produces 1,152 scenarios per reference graph before the 12 probe settings. Named-storage functions are varied only where the reports identify them. Equal scenario-cell weighting is a design convention, not an estimate of how often those conditions occur in nature.')
table(['Module','Routing evaluations / scope'],[
('Six neutral reference models + two storage-sensitive models','11,520 routing evaluations.'),
('Area-weight sensitivity','1,728 evaluations; Kuranji, Huong and Kadamaian only.'),
('CDO matched rewiring','17,280 evaluations: 15 labelled four-source binary arrangements, including the reference.'),
('Kadamaian matched rewiring','3,456 evaluations: all three labelled three-source binary arrangements.'),
('Huong branch/target placement','2,304 evaluations: Bo-before-Hue counterfactual and alternative residual entry.'),
('Kuranji order uncertainty','7,560 evaluations: 105 labelled binary refinements × 72 prespecified subset scenarios.'),
('Strictly paired timing contrast','6,144 evaluations: source starts changed, all other inputs fixed.'),
('Outlet failure/common-boundary contrast','82,944 extra probe evaluations; 1,152 previously defined upstream settings recomputed.')],[2.5,4.1])
p('The main/control modules total 43,848 routes; adding paired timing gives 49,992. The outlet module reuses existing reference settings and does not add independent source scenarios. Rewiring preserves source input; it is not removal of a tributary’s water.')

page('4. Primary results at the receiving probes')
p('The table reports the neutral-special-storage baseline, not present-day reservoir operation. Ordinary link memory still varies according to the common scenario grammar. Every receiving probe receives the full normalized event volume. These are descriptive outcomes of coarse models, not a ranking of actual city flood risk.')
rows=[]
for case,label in [('CDO','CDO lower corridor'),('MANADO','Lower Manado'),('KURANJI','Lower Kuranji'),('HUE','Huong lagoon aggregate'),('KADAMAIAN','Kadamaian lower aggregate'),('HATYAI','Songkhla receiving aggregate')]:
 z=main[(main.case==case)&main.role.isin(['receiver','urban_and_receiver'])].iloc[0]
 rows.append([label,'13,824',pct(z.mean_spill_fraction),pct(z.overflow_share)])
table(['Probe','Evaluations','Mean input spilled','Cases with any spill'],rows,[2.3,1.,1.55,1.65])
p('“Mean input spilled” is the arithmetic mean of spill/100. “Cases with any spill” is the percentage of tested parameter settings with spill > 10⁻⁸ units. Zero-spill cases are retained in both summaries. The denominator is not a count of observed storms.')
h('Target position changes what reaches the target',2)
z=main[(main.case=='HUE')&(main.role=='urban')].iloc[0]
p(f'The pre-Sinh Hue/Kim Long probe had mean spill {pct(z.mean_spill_fraction)}, versus 17.18% at the lagoon receiving probe. This is not an independent comparison of two cities: the upstream probe does not directly receive Bo or the provisionally downstream residual source. Its spill is still normalized by total basin input of 100 units. The lower value must not be read as measured local safety.')
p('The Hat Yai urban entry probe had 21.01% mean spill; its lake-receiver probe had 17.95%. In the forward routing model, additional link travel and memory can spread arrivals. No result here establishes that lake-side flooding is milder than city flooding in reality.')
h('What the table can and cannot establish',2)
p('All models have explicit source convergence, but their named source groups are at different levels of aggregation. A source group may summarize an entire subcatchment. Therefore source count and graph-hop count are not standardized high-resolution drainage metrics. Differences in these baseline means cannot be attributed solely to a verified geographic motif.')
p('The causal evidence inside the computational model comes from paired interventions that hold input and resource budgets fixed, not from the cross-case means alone. Six selected analogue reports are not a random sample of Southeast Asian rivers or a flood/non-flood case-control design.')
p('Data location: results/summary_primary.csv; full probe records: results/probe_results.csv.gz. Structural descriptors and source-to-receiver graph hops are in results/topology_descriptors.csv.')

page('5. Paired timing: a recurring but conditional effect')
fig('01_paired_timing','Figure 1. Change only source starting times, preserving source weights, pulse shapes, active sources, link delays, storage functions, capacity and buffer. Positive bars indicate more mean spill with simultaneous source starts.')
rows=[]
for case in ['CDO','MANADO','KURANJI','HUE','KADAMAIAN','HATYAI']:
 z=timing[(timing.case==case)&(timing.storage_tau==0)].iloc[0]
 rows.append([case,pp(z.mean_delta),pct(z.higher_share),pct(z.lower_share),pct(z.tie_share)])
table(['Model','Mean change','Higher spill','Lower spill','Tied'],rows,[1.5,1.35,1.25,1.25,1.15])
p('Each row summarizes 4,608 matched probe-setting pairs. In all six models the ensemble-mean direction is positive. Nevertheless, CDO has a decrease in 4.67% of pairs; Manado in 5.47%; Kadamaian in 5.69%. Equal source starting times do not guarantee equal arrival times after unequal routing paths.')
p('With strong named storage enabled, the mean timing effect remained positive but fell to +2.42 percentage points for Manado and +1.86 for Huong. Adverse retiming remains possible. This is evidence for a conditional arrival-overlap mechanism, not an inevitable-failure rule.')
p('Changing source timing is a forcing intervention, not a topology change. It tests the response of a fixed topology. The next section changes the structure itself. No correlation of a peak-based synchrony index with peak overflow is presented as independent causal proof.')

page('6. CDO structural intervention: retain sources, change connections')
fig('02_cdo_rewiring','Figure 2. Four source groups and three binary merging nodes are retained. There are 15 labelled arrangements but only two unlabelled shapes: balanced and sequential. ID 12 reproduces the report-derived CDO order exactly.')
p('The supplied coarse order is ((Upper CDO + Tumalaong) + Bubunawan) + lower laterals. Its mean spill was 18.39%. The three balanced arrangements averaged 20.49%, an increase of 2.10 percentage points. This demonstrates that a connection change can alter the response under fixed event volumes and graph-size budgets, but it does not show that CDO’s order is uniquely hazardous.')
f=findings['CDO']
p(f"Across all matched settings, the reference’s median rank was {f['median_rank']:.1f}/15, with ties averaged and rank 1 denoting least spill. Of {f['total_settings']:,} settings, {f['informative_settings']:,} distinguished at least two arrangements; among those informative settings the median rank was {f['median_informative_rank']:.1f}/15. Ranking the ensemble means instead places the reference {int(f['rank_of_ensemble_mean'])}/15. These are different summaries and must not be conflated.")
p('Balanced alternatives produced more spill than the CDO order in 37.27% of paired comparisons, less spill in 13.16%, and a tie in 49.57%. An arrangement that is worse on average can therefore be better for a particular loading pattern.')
h('Recorded counterexample',2)
p('CDO scenario 121, capacity 20, buffer 0: the reference spilled 45.71% of input, while balanced alternative 04 spilled 0.14%. Inputs, source activation, durations and capacity were retained. This case directly rejects an unconditional “balanced merging always spills more” statement.')
p('The identity check between the reference model and alternative 12 produced zero output difference. All raw alternatives, tree strings, matched outcomes and scenario IDs are archived. These hypothetical rearrangements are not proposed river construction works.')

page('7. Named storage functions modify the structural response')
fig('03_storage_effect','Figure 3. Connections remain fixed while conservative storage-release functions are changed. Release-time parameters are dimensionless hypotheses, not measured reservoir operations.')
table(['Receiver','Neutral storage','Strong routing storage','Paired mean change'],[
('Manado','22.05%','14.08%','−7.97 percentage points'),
('Huong lagoon','17.18%','7.30%','−9.88 percentage points'),
('Hue/Kim Long','5.46%','0.95%','−4.52 percentage points')],[2.1,1.3,1.45,1.65])
p('These reductions do not mean that the named dams actually achieve these percentages. The reports identify where stateful storage/control nodes belong; this experiment assigns generic release functions at those locations. No reservoir rule curves, initial lake levels or hydropower dispatch records were available or used. [D2, pp. 5, 12; D4, pp. 11–12.]')
h('Storage is not universally protective in the tested ensemble',2)
p('Manado scenario 529 uses an eight-tick block pulse, equal active-source weights, a two-of-three active-source subset, unit link delays and staggered starts. At capacity 5 and buffer 10, the neutral-control model spills 0.00%, while the strong-storage model spills 17.41%. Routing storage retimes one contribution into the arrival window of another. The complete source-tagged calculation is reproducible; Figure 6 in the archive illustrates this counterexample.')
p('Strong storage increased spill in 2.29% of all Manado probe settings and decreased it in 42.03%; the rest tied. For Huong’s lagoon receiver, increases occurred in 0.47% and decreases in 46.74%. Neutral and storage-modified graphs therefore need separate interpretation.')
p('A topology-focused scope may exclude measured infrastructure performance as an input requirement. It cannot infer that capacity, storage or operation are irrelevant: these controlled results show that changing such functions on a fixed graph changes the response.')

page('8. Incomplete order and urban position affect conclusions')
fig('04_kuranji_uncertainty','Figure 4. Every refinement preserves the five Kuranji source groups and total input. The 105 labelled binary arrangements span three unlabelled shapes. They are an uncertainty envelope, not 105 verified river networks.')
p('On the same 72 selected source/routing scenarios and 12 probe settings, the collapsed middle-zone Kuranji graph gave 24.18% mean spill. Alternative binary orders gave 19.31–23.03%, with a median of 20.23%. Adding unresolved merger structure adds logical routing steps; the result depends on that modelling choice. The collapsed zone should not be described as a surveyed five-way confluence. [D3, pp. 1–3, 7, 13.]')
h('Huong demonstrates why the target must be located in the graph',2)
p('The reference preserves Tuan → Hue/Kim Long → Sinh, with Bo joining at Sinh. Moving Bo to Tuan in a counterfactual raised mean spill at the Hue probe from 5.46% to 12.01%, while the lagoon mean changed from 17.18% to 17.37%. Total basin input stayed at 100 units; the intervention changed which sources pass the urban target. [D4, pp. 3, 10.]')
p('The 445 km² residual floodplain component is not located precisely by a local source polygon in the supplied report. Moving its coarse injection from Sinh to Hue, as a separate uncertainty case, raised Hue-probe mean spill to 11.39%. This is a boundary-assignment sensitivity, not evidence that all residual runoff actually enters Hue at that node.')
h('Area-weight sensitivity is separate from connectivity',2)
p('Using the supplied internally consistent source partitions instead of equal weights increased mean receiver spill by 1.59 percentage points in Kuranji, 0.16 in Huong and 0.16 in Kadamaian on matched equal-weight settings. This assumes equal runoff per unit area for that sensitivity only; it does not turn area into observed discharge. [D3, p. 7; D4, p. 6; D5, p. 11.]')

page('9. Outlet routes and shared downstream dependence')
fig('05_shared_boundary','Figure 5. One-route and two-route buffers receive the same incoming series and total capacity/storage budget. The two-route case is not granted double capacity. This is an abstract service/failure experiment.')
p('At Hat Yai’s urban receiving probe, unrestricted mean spill was 21.01% for both one and two equal-budget routes. Reducing total effective service by 50% for a common downstream condition raised both to 35.22%. Merely drawing a second route did not help when the same total capacity and common multiplier were retained.')
p('When one route was completely blocked, the single-route model lost all service. The two-route model retained half of its original total service and had 35.22% mean spill. In the zero-service single-route case, mean spill was 86.67% and the remaining input stayed in the buffer; total unreleased input was 100%. Stored water is accounted for, not silently deleted.')
h('The structural part of the result',2)
p('In the report-derived Hat Yai coarse graph, two edge-disjoint paths connect the urban zone to the lake aggregate, but only one aggregate edge connects the common lake node to the final Gulf boundary. Computed minimum edge-cut sizes are therefore 2 and 1, respectively. Removing the shared downstream connection defeats both routes. This is a property of the coarse graph, not a surveyed count of every actual outlet. [D6, pp. 10, 21, 24–26.]')
p('Capacity multipliers stand for a generic common service limitation. They were not derived from a real lake level and do not solve backwater. A forward-only graph cannot establish which neighbourhood accumulates water when the lake rises.')
p('The general implication is conditional redundancy: alternative paths protect against a failure that they do not share. They do not automatically remove a common downstream dependency. Many single-outlet graphs share this feature, so it cannot identify CDO as exceptional by itself.')

page('10. Interpretation for a Lapasan-motivated thesis')
h('What the evidence supports',2)
p('Within this model class, source-path convergence, the position of merging branches relative to the target, and the placement of storage or common receiving constraints influence normalized overflow. Paired timing experiments provide a consistent positive mean response across six coarse systems; source-preserving CDO rewiring demonstrates an actual structural effect inside the model. Counterexamples show that neither the sign of an individual comparison nor a site ranking is invariant.')
h('What the evidence does not establish',2)
p('The study does not establish that the Kitanglad-facing CDO network causes the observed floods in Barangay Lapasan, that local drainage or flood controls are irrelevant, that every nearby lowland is a receiving node of the CDO mainstem, or that flooding is inevitable. The CDO report’s mouth and corridor context supports a regional hypothesis; map proximity alone does not supply an edge, overbank pathway or hydraulic exchange to Lapasan. [D1, pp. 4, 7–9.]')
p('The primary output should therefore be described as a document-constrained, comparative hydro-topological stress experiment motivated by Lapasan, not a validated Lapasan flood simulation. The graph labelled CDO is not renamed Lapasan in the data or figures. Determining local relevance requires evidence of the relevant connection or mechanism, but it does not require expanding this work into a fully government-data-dependent forecasting project.')
h('Suggested results/conclusion paragraph',2)
p('“A comparative, dimensionless routing experiment was conducted using six coarse networks constrained by supplied regional river-system reports. Under matched source volumes and routing assumptions, simultaneous source starts increased ensemble-mean receiving-node spill in all six models. Source-preserving rewiring of the CDO model changed its response, while storage-function, confluence-order and shared-outlet tests revealed substantial conditionality. The findings support arrival concentration and downstream dependency as candidate structural mechanisms, but do not establish an inevitable or uniquely CDO-specific flood pattern. Their relevance to Barangay Lapasan remains a regional mechanistic hypothesis rather than a validated local inundation result.”')
h('Publication and inference limits',2)
p('These results can support a transparent computational methods/results section, subject to adviser review and source verification. Publication, acceptance and novelty are not established by this report. There are six selected report cases, not hundreds of thousands of independent observations. No p-values, real-event probabilities or fitted claims of geographical predictive accuracy are supplied. Present the original-source verification status and aggregation uncertainty explicitly.')
p('Do not reuse the earlier generic synthetic pilot as geographic validation. This execution is a new common-model experiment tied to the six uploaded reports; its synthetic parameter choices remain clearly identified.')

page('11. Verification, source keys and reproduction')
table(['Verification','Executed result'],[
('Core unit/property checks','26 passed: source consistency, acyclicity, split conservation, hand queue calculation, monotonicity, source ordering and graph relabel checks.'),
('Whole-run water accounting','Maximum absolute routing residual error < 1 × 10⁻¹² units; maximum probe water-accounting error < 1 × 10⁻¹² units.'),
('Finite-horizon checks','Selected 512- versus 1,024-tick comparisons agree within 1 × 10⁻⁷ units. Main probe buffers fully drain at positive capacities.'),
('Exact reference identity','CDO alternative 12 exactly matches the reference output; maximum spill-fraction difference = 0.'),
('Equal-budget outlet null test','One/two route outcomes exactly match under unrestricted and shared-50% conditions.'),
('Time-bin subdivision','Selected constant-rate buffer bins divided into halves agree within 6 × 10⁻¹³ units. This is not full hydraulic timestep convergence.'),
('Not performed','Field calibration; validation against independent flood events; surveyed geometry verification; inundation-extent or flood-depth skill assessment.')],[2.2,4.4])
h('Input report keys',2)
p('D1 is dated 29 August 2026; D2–D6 are dated 11 September 2026. Full filenames and SHA-256 fingerprints appear in inputs/source_manifest.json; these dossiers were not independently peer reviewed or source-verified in this execution.')
for i,sp in enumerate(reference_specs(),1):
 p(f'D{i}. {sp.label}. User-supplied deep terrain/hydrology/topology technical report.')
h('How to reproduce',2)
p('Install the versions in requirements.txt, then run python run_all.py from the extracted package directory. Tests run first, then the main experiment, paired timing and analysis. No additional data download is required for these coarse models. To reproduce only the checks: python code/test_model.py.')
p('Inspect inputs/*_graph.json and inputs/edge_evidence_register.csv before interpreting outputs. Every scenario has a scenario_id. Main route records preserve actual weights, starting times, active sources and water balance. Full probe records use gzip-compressed CSV. The HTML viewer displays a selected equal-input/unit-delay subset, not the entire factor grid.')
p('SI is the generalized source-tag synchrony index: combined peak divided by the sum of each source-tagged peak at that probe. It is an output diagnostic; it is not independent evidence that the same ratio predicts its own numerator. No hidden drains, local flood footprints or government-restricted measurements are assumed.')

# Six appendix pages: a visible audit trail from sources to model, not invented geography.
for i,sp in enumerate(reference_specs(),1):
 page(f'Appendix {chr(64+i)}. {sp.label}')
 fig('graph_'+sp.case,f'Graph {i}. Exact coarse routing structure used in the reference configuration. Labels come from D{i}; geometric positions on this drawing are arbitrary.')
 p(sp.provenance)
 h('Preserved information and unresolved detail',2);p(sp.limitations)
 met=desc[desc.case==sp.case].iloc[0]
 table(['Coarse descriptor','Value'],[
 ('Represented source groups',str(met.source_groups)),('Model vertices / edges',f'{met.nodes} / {met.edges}'),
 ('Source-to-receiver logical hops',str(met.receiver_hops)),('Direct source-group fraction upstream of urban probe',f'{100*met.urban_source_fraction:.0f}% (unweighted groups, not measured runoff)'),
 ('Receiver node / urban node',f"{sp.observations['receiver']} / {sp.observations['urban']}"),
 ('Status of named control functions','Synthetic conservative release times; zero in neutral runs' if sp.storage else 'No extra named upstream storage function assigned')],[3.2,3.4])
 if sp.areas is not None:
  h('Separate source-area sensitivity',2)
  p('; '.join(f'{sp.names[n]}: {ar:g} km²' for n,ar in zip(sp.sources,sp.areas))+'. Values are used only as normalized weights under an equal-runoff-per-area assumption. No alternative incompatible basin total is averaged into this partition.')
 else:
  p('No complete, internally consistent source-area partition was assigned. Equal and randomized weights are synthetic scenario inputs, not claimed observed subcatchment contributions.')
 p('Reproduction files: inputs/'+sp.case+'_graph.json; inputs/node_register.csv; inputs/edge_evidence_register.csv; results/probe_results.csv.gz. A GIS extraction threshold, raw drainage-line inventory and geographical channel coordinates are not available in this model.')

# Metadata without fabricated author affiliation or publication status.
doc.core_properties.title='Six-Case Hydro-Topological Stress Simulation'
doc.core_properties.subject='Document-constrained coarse-network experiment motivated by Lapasan'
doc.core_properties.author=''
doc.core_properties.comments='Executed computational supplement. Synthetic parameters; no mapped inundation forecast.'
outfile=root.parent/'SEA_Topology_Executed_Methods_and_Results.docx';doc.save(outfile)
print(outfile)
